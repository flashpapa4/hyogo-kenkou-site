# -*- coding: utf-8 -*-
"""
兵庫県 令和8年度健康づくり研修会支援事業
講師育成研修（90分）で使用する配布シート3種を生成する。

  1. 講座仕分けシート（受講者記入・ブロック2）
  2. 担当講座確定シート（受講者記入・ブロック6）
  3. 実施記録シート（育成担当者記入・8名分の記録用）

出力: docs/講師育成研修_90分_配布シート.docx
実行: python3 docs/tools/build-instructor-training-worksheets.py [出力パス]
"""
import sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

DEEP = RGBColor(0x1B, 0x43, 0x32)
GREEN = RGBColor(0x2D, 0x6A, 0x4F)
SUB = RGBColor(0x63, 0x6E, 0x72)
JP = "Meiryo"

COURSES = [
    ("ア. 運動習慣の定着", [
        ("1", "デスクワーク疲労を90分でリセット", "90分／対面・オンライン"),
        ("2", "見た目年齢マイナス5歳！姿勢改善プログラム", "90分／対面・オンライン"),
        ("3", "昼休み15分でルーティン定着（5日間チャレンジ）", "15分×5日／オンライン限定"),
    ]),
    ("イ. 食生活の改善", [
        ("4", "健診で引っかからないカラダの作り方（食事編）", "90分／対面・オンライン"),
        ("5", "もう失敗しないダイエットの成功法則", "90分／対面・オンライン"),
    ]),
    ("ウ. 健康意識向上・行動変容促進", [
        ("6", "まずは知る、カラダ年齢（体力測定）", "1人30分／対面のみ"),
        ("7", "見えない数値を見える化（体成分測定）", "1人30分／対面のみ"),
        ("8", "健診で血圧を指摘された方へ", "90分／対面・オンライン"),
        ("9", "対人ストレスに効く カラダから整えるメンタルケア", "90分／対面・オンライン"),
        ("10", "PMS・更年期に対応する 女性社員のためのセルフケア", "90分／対面・オンライン"),
    ]),
]


def set_font(run, size=10.5, bold=False, color=None):
    run.font.name = JP
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color
    run._element.rPr.rFonts.set(qn("w:eastAsia"), JP)


def para(doc, text, size=10.5, bold=False, color=None, align=None, space_after=4):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    set_font(p.add_run(text), size, bold, color)
    return p


def sheet_header(doc, no, title, who):
    para(doc, "兵庫県 令和8年度 健康づくり研修会支援事業　講師育成研修", 9, color=SUB, space_after=2)
    para(doc, f"シート{no}　{title}", 17, bold=True, color=DEEP, space_after=2)
    para(doc, who, 9.5, color=GREEN, space_after=8)


def style_table(table, widths):
    """列幅を固定する。python-docx の既定は自動レイアウトのため、
    tblLayout を fixed にし、tblGrid と各セルの幅を明示しないと列幅が反映されない。"""
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tblPr = table._tbl.tblPr
    layout = tblPr.makeelement(qn("w:tblLayout"), {qn("w:type"): "fixed"})
    tblPr.append(layout)
    grid = table._tbl.find(qn("w:tblGrid"))
    if grid is not None:
        for col, cm in zip(grid.findall(qn("w:gridCol")), widths):
            col.set(qn("w:w"), str(int(Cm(cm).twips)))
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            if i < len(widths):
                cell.width = Cm(widths[i])
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.space_before = Pt(2)


def cell_text(cell, text, size=10, bold=False, color=None, align=None):
    p = cell.paragraphs[0]
    if align is not None:
        p.alignment = align
    set_font(p.add_run(text), size, bold, color)


doc = Document()
sec = doc.sections[0]
sec.orientation = WD_ORIENT.PORTRAIT
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
for attr, v in (("top_margin", 1.5), ("bottom_margin", 1.5), ("left_margin", 1.6), ("right_margin", 1.6)):
    setattr(sec, attr, Cm(v))
style = doc.styles["Normal"]
style.font.name = JP
style.font.size = Pt(10.5)
style.element.rPr.rFonts.set(qn("w:eastAsia"), JP)

# ---------------- シート1：講座仕分けシート ----------------
sheet_header(doc, "1", "講座仕分けシート", "記入者：受講者　／　記入タイミング：研修中（10講座の紹介のあと）")
para(doc, "10講座すべてに ◎ △ ✕ のいずれかを記入してください。迷ったら △ で構いません。", 10, space_after=2)
para(doc, "◎ できる　　△ 要相談（準備・確認が必要）　　✕ 今回は外す", 9.5, color=GREEN, space_after=8)

t = doc.add_table(rows=1, cols=5)
t.style = "Table Grid"
for i, h in enumerate(["", "講座名", "時間・形式", "◎△✕", "気になっている点・確認したいこと"]):
    cell_text(t.rows[0].cells[i], h, 9.5, bold=True, color=DEEP,
              align=WD_ALIGN_PARAGRAPH.CENTER if i in (0, 3) else None)
for cat, rows in COURSES:
    r = t.add_row()
    r.cells[0].merge(r.cells[4])
    cell_text(r.cells[0], cat, 10, bold=True, color=GREEN)
    for no, name, fmt in rows:
        r = t.add_row()
        cell_text(r.cells[0], no, 10, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        cell_text(r.cells[1], name, 9.5)
        cell_text(r.cells[2], fmt, 9)
        cell_text(r.cells[3], "", 10)
        cell_text(r.cells[4], "", 9)
style_table(t, [1.0, 6.4, 3.6, 1.6, 5.2])
para(doc, "", 6, space_after=2)
para(doc, "「実技はできるが説明が不安」という理由で ✕ にしないでください。説明はスライドと講師台本でカバーできます。",
     9.5, color=SUB, space_after=0)

# ---------------- シート2：担当講座確定シート ----------------
doc.add_page_break()
sheet_header(doc, "2", "担当講座確定シート", "記入者：受講者　／　記入タイミング：研修の最後（ブロック6）")

t = doc.add_table(rows=2, cols=4)
t.style = "Table Grid"
for i, h in enumerate(["お名前", "実施日", "対応可能な形式", "連絡先"]):
    cell_text(t.rows[0].cells[i], h, 9.5, bold=True, color=DEEP)
for i, v in enumerate(["", "　　　年　　月　　日", "対面 ・ オンライン ・ 両方", ""]):
    cell_text(t.rows[1].cells[i], v, 10)
style_table(t, [4.4, 4.4, 4.6, 4.4])

para(doc, "", 8, space_after=2)
para(doc, "担当する講座", 12, bold=True, color=DEEP, space_after=4)
t = doc.add_table(rows=1, cols=4)
t.style = "Table Grid"
for i, h in enumerate(["講座番号", "講座名", "資料受領", "備考（準備が必要なこと）"]):
    cell_text(t.rows[0].cells[i], h, 9.5, bold=True, color=DEEP,
              align=WD_ALIGN_PARAGRAPH.CENTER if i in (0, 2) else None)
for _ in range(6):
    r = t.add_row()
    for i in range(4):
        cell_text(r.cells[i], "", 10)
style_table(t, [2.2, 7.4, 2.2, 6.0])

para(doc, "", 8, space_after=2)
para(doc, "確認事項", 12, bold=True, color=DEEP, space_after=4)
for line in [
    "□　上記の講座について、研修当日までに資料を読み込みます",
    "□　受講料・相談料等を受け取らないこと、自社の営業をしないことを理解しました",
    "□　実技の種目は自分で組み立ててよいこと、説明パートは台本どおりでよいことを理解しました",
    "□　あとから担当講座を追加・変更できることを確認しました",
]:
    para(doc, line, 10.5, space_after=5)

para(doc, "", 10, space_after=4)
para(doc, "次回チェック日：　　　　年　　月　　日　　　時　　分　〜　（60分）　　　場所・方法：", 10.5, bold=True, space_after=6)
para(doc, "※ このチェックは合否を決めるものではありません。研修当日までに足りないものを一緒に確認し、必要なサポートを決めるためのものです。",
     9.5, color=SUB, space_after=6)
para(doc, "ご署名　　　　　　　　　　　　　　　　　　　　　　　担当者確認　　　　　　　　　　　　　　　", 10.5, space_after=0)

# ---------------- シート3：実施記録シート ----------------
doc.add_page_break()
sheet_header(doc, "3", "実施記録シート", "記入者：育成担当者　／　1名につき1枚。60分チェックにそのまま引き継ぐ")

t = doc.add_table(rows=2, cols=4)
t.style = "Table Grid"
for i, h in enumerate(["受講者名", "実施日", "保有資格", "普段の指導内容・得意分野"]):
    cell_text(t.rows[0].cells[i], h, 9.5, bold=True, color=DEEP)
for i in range(4):
    cell_text(t.rows[1].cells[i], "", 10)
style_table(t, [3.4, 3.4, 4.2, 6.8])

para(doc, "", 8, space_after=2)
blocks = [
    ("確認クイズ（5問）", "間違えた問題番号：　　　　　　　　　　　　　　　　　　　　　→ 60分チェックで重点確認"),
    ("説明パートの実演", "良かった点：\n\n改善を伝えた点：\n\nリテイクでの改善："),
    ("実技パートの実演", "得意分野・実施した種目：\n\n体調確認の声かけ　有 ・ 無　／　座位の代替種目の提示　有 ・ 無　／　中止指示　有 ・ 無\n\n伝えた点："),
    ("つなぎの練習", "説明→実技　／　実技→説明　／　実技→習慣化シート　　　できた項目に〇："),
    ("安全・実務のケース問答", "Q1　可 ・ 要補足　　Q2　可 ・ 要補足　　Q3　可 ・ 要補足"),
    ("60分チェックまでの補強点", ""),
]
for title, hint in blocks:
    para(doc, title, 11, bold=True, color=GREEN, space_after=3)
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    c = t.rows[0].cells[0]
    cell_text(c, hint, 9.5, color=SUB)
    for _ in range(2 if hint.count("\n") < 2 else 3):
        set_font(c.add_paragraph().add_run(""), 9.5)
    style_table(t, [17.8])
    para(doc, "", 5, space_after=0)

out = sys.argv[1] if len(sys.argv) > 1 else "docs/講師育成研修_90分_配布シート.docx"
doc.save(out)
print("written:", out)
