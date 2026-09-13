# -*- coding: utf-8 -*-
"""
兵庫県 令和8年度健康づくり研修会支援事業
講師育成 60分チェック（実施前チェック）で使用する資料一式を生成する。

  1. 実施要領（育成担当者用・60分進行表）
  2. 確認クイズ〇×5問＋口頭確認（受講者記入）
  3. 実演評価シート（説明／実技／つなぎ）
  4. ケース問答（育成担当者用・模範初動つき）
  5. 結果と補強プラン（受講者控え・A/B/C）

このチェックは合否判定ではなく、研修当日までの過不足を確認し、
必要なサポートを決めるためのもの。「担当させない」という結果は作らない。

出力: docs/講師育成_60分チェック_実施セット.docx
実行: python3 docs/tools/build-check-materials.py [出力パス]
"""
import sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

DEEP = RGBColor(0x1B, 0x43, 0x32)
GREEN = RGBColor(0x2D, 0x6A, 0x4F)
ACCENT = RGBColor(0xE7, 0x6F, 0x51)
SUB = RGBColor(0x63, 0x6E, 0x72)
JP = "Meiryo"


def set_font(run, size=10.5, bold=False, color=None):
    run.font.name = JP
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color
    run._element.rPr.rFonts.set(qn("w:eastAsia"), JP)


def para(doc, text, size=10.5, bold=False, color=None, align=None, space_after=4):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    set_font(p.add_run(text), size, bold, color)
    return p


def sheet_header(doc, no, title, who):
    para(doc, "兵庫県 令和8年度 健康づくり研修会支援事業　講師育成 60分チェック", 9, color=SUB, space_after=2)
    para(doc, f"シート{no}　{title}", 17, bold=True, color=DEEP, space_after=2)
    para(doc, who, 9.5, color=GREEN, space_after=8)


def style_table(table, widths):
    """列幅を固定する（python-docx の既定は自動レイアウトのため tblLayout を fixed にする）。"""
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tblPr = table._tbl.tblPr
    tblPr.append(tblPr.makeelement(qn("w:tblLayout"), {qn("w:type"): "fixed"}))
    grid = table._tbl.find(qn("w:tblGrid"))
    if grid is not None:
        for col, cm in zip(grid.findall(qn("w:gridCol")), widths):
            col.set(qn("w:w"), str(int(Cm(cm).twips)))
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            if i < len(widths):
                cell.width = Cm(widths[i])
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.space_before = Pt(2)


def cell_text(cell, text, size=10, bold=False, color=None, align=None):
    p = cell.paragraphs[0]
    if align is not None:
        p.alignment = align
    set_font(p.add_run(text), size, bold, color)


def header_row(table, labels, widths, center_idx=()):
    for i, h in enumerate(labels):
        cell_text(table.rows[0].cells[i], h, 9.5, bold=True, color=DEEP,
                  align=WD_ALIGN_PARAGRAPH.CENTER if i in center_idx else None)


def blank_box(doc, hint, lines=2, width=17.8):
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    c = t.rows[0].cells[0]
    cell_text(c, hint, 9.5, color=SUB)
    for _ in range(lines):
        set_font(c.add_paragraph().add_run(""), 9.5)
    style_table(t, [width])
    para(doc, "", 5, space_after=0)
    return t


doc = Document()
sec = doc.sections[0]
sec.orientation = WD_ORIENT.PORTRAIT
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
for attr, v in (("top_margin", 1.5), ("bottom_margin", 1.5), ("left_margin", 1.6), ("right_margin", 1.6)):
    setattr(sec, attr, Cm(v))
style = doc.styles["Normal"]
style.font.name = JP
style.font.size = Pt(10.5)
style.element.rPr.rFonts.set(qn("w:eastAsia"), JP)

# ---------------- シート1：実施要領 ----------------
sheet_header(doc, "1", "実施要領（60分）", "使用者：育成担当者　／　1名につき60分・個別実施")

para(doc, "このチェックの位置づけ", 12, bold=True, color=DEEP, space_after=3)
for line in [
    "合否を決めるものではありません。研修当日までに足りないものを確認し、必要なサポートを決めるためのものです。",
    "「担当していただかない」という結果は作りません。結果は A／B／C のいずれかで、どう支えれば担当できるかを示します。",
    "見る観点は90分研修で開示済みです。隠して測るのではなく、事前に伝えたとおりに確認します。",
]:
    para(doc, "・" + line, 10, space_after=3)

para(doc, "", 6, space_after=2)
para(doc, "開始時に必ず伝える3つ", 12, bold=True, color=DEEP, space_after=3)
for line in [
    "これは合否判定ではありません。今日の目的は、当日までに準備するものをはっきりさせることです。",
    "見るところは研修でお伝えした4点（ルールの理解／説明パート／実技パート／安全への対応）だけです。",
    "できていない点があっても、こちらで資料やサポートを用意します。遠慮なく分からないと言ってください。",
]:
    para(doc, "・" + line, 10, space_after=3)

para(doc, "", 6, space_after=2)
para(doc, "進行表", 12, bold=True, color=DEEP, space_after=4)
t = doc.add_table(rows=1, cols=4)
t.style = "Table Grid"
header_row(t, ["時間", "内容", "使用シート", "確認するポイント"], [1.8, 4.6, 2.6, 8.8], center_idx=(0,))
for row in [
    ("10分", "ルール・安全の口頭確認", "シート2", "禁止事項3点／アンケート／迷ったときの連絡先"),
    ("15分", "説明パートの実演", "シート3", "担当講座から1パート。資料を見ながらで可"),
    ("10分", "実技パートの実演", "シート3", "得意分野で可。安全の声かけと代替種目"),
    ("10分", "ケース問答", "シート4", "初動が言えるか。正解の暗記は不要"),
    ("10分", "フィードバックと補強プラン", "シート5", "良い点を先に。A／B／Cをその場で伝える"),
    ("5分", "実施スケジュール・事務連絡", "—", "初回の日程、同行の要否、連絡方法"),
]:
    r = t.add_row()
    cell_text(r.cells[0], row[0], 10, bold=True, color=GREEN, align=WD_ALIGN_PARAGRAPH.CENTER)
    cell_text(r.cells[1], row[1], 10, bold=True)
    cell_text(r.cells[2], row[2], 9.5)
    cell_text(r.cells[3], row[3], 9.5)
style_table(t, [1.8, 4.6, 2.6, 8.8])

para(doc, "", 8, space_after=2)
para(doc, "進め方の注意", 12, bold=True, color=DEEP, space_after=3)
for line in [
    "良かった点を先に、具体的に伝える。改善点は3つまで。多く言うと当日までに消化できません。",
    "実演の途中で口を挟まない。止めるのは安全上の問題があるときだけ。",
    "A／B／C はその場で本人に伝え、理由を1文で言う。持ち帰って後日通知にしない。",
    "8名とも同じ設問・同じ観点で実施する。実施者が複数の場合はシート3・4をそのまま使う。",
    "記録は当日中にシート5の控えを本人に渡し、原本を事務局で保管する。",
]:
    para(doc, "・" + line, 10, space_after=3)

# ---------------- シート2：確認クイズ ----------------
doc.add_page_break()
sheet_header(doc, "2", "確認クイズ（〇×5問）と口頭確認", "記入者：受講者　／　10分　※90分研修と同一の設問です")

para(doc, "次の内容が正しければ〇、誤っていれば×を記入してください。", 10, space_after=6)
t = doc.add_table(rows=1, cols=3)
t.style = "Table Grid"
header_row(t, ["", "設問", "〇／×"], [1.0, 14.0, 2.8], center_idx=(0, 2))
for i, q in enumerate([
    "受講者から実費程度であればテキスト代を受け取ってもよい",
    "研修の最後に自社のサービス案内を配ってもよい",
    "1社あたりの利用は令和6年度以降の通算で3回までである",
    "遠方の企業からの依頼は距離を理由に断ることができる",
    "研修終了後には個別質疑応答の時間を設ける必要がある",
], 1):
    r = t.add_row()
    cell_text(r.cells[0], str(i), 10, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    cell_text(r.cells[1], q, 10)
    cell_text(r.cells[2], "", 10)
style_table(t, [1.0, 14.0, 2.8])
para(doc, "（正解：1 ×　2 ×　3 〇　4 ×　5 〇　※採点後に本人へ開示する）", 9, color=SUB, space_after=8)

para(doc, "口頭で確認すること", 12, bold=True, color=DEEP, space_after=4)
t = doc.add_table(rows=1, cols=3)
t.style = "Table Grid"
header_row(t, ["質問", "確認したい答え", "可／要補足"], [6.4, 8.6, 2.8], center_idx=(2,))
for q, a in [
    ("アンケートはいつ、何回とりますか", "研修前・研修直後・1ヶ月後。終了前にその場で回収する"),
    ("この事業の成果指標は何ですか", "週1回以上運動する人の割合が受講者全体で60%以上"),
    ("判断に迷ったときはどうしますか", "その場で答えず「確認してご連絡します」。事務局へ連絡する"),
    ("受講者に自社の施設を尋ねられたら", "案内・勧誘はしない。営業・広告宣伝は禁止されている"),
]:
    r = t.add_row()
    cell_text(r.cells[0], q, 9.5)
    cell_text(r.cells[1], a, 9.5, color=SUB)
    cell_text(r.cells[2], "", 10)
style_table(t, [6.4, 8.6, 2.8])
para(doc, "", 6, space_after=2)
para(doc, "※ 間違えた項目はその場で正解を伝え、シート5の「補強する点」に記入する。", 9.5, color=SUB, space_after=0)

# ---------------- シート3：実演評価シート ----------------
doc.add_page_break()
sheet_header(doc, "3", "実演評価シート", "記入者：育成担当者　／　説明15分・実技10分")

hdr = doc.add_table(rows=2, cols=4)
hdr.style = "Table Grid"
header_row(hdr, ["受講者名", "実施日", "対象講座", "形式"], [3.6, 4.4, 6.2, 3.6])
for i, v in enumerate(["", "　　　年　　月　　日", "", "対面 ・ オンライン"]):
    cell_text(hdr.rows[1].cells[i], v, 10)
style_table(hdr, [3.6, 4.4, 6.2, 3.6])
para(doc, "", 8, space_after=2)

def check_table(title, items):
    para(doc, title, 12, bold=True, color=GREEN, space_after=4)
    t = doc.add_table(rows=1, cols=3)
    t.style = "Table Grid"
    header_row(t, ["観点", "可／要補足", "メモ"], [8.0, 2.4, 7.4], center_idx=(1,))
    for it in items:
        r = t.add_row()
        cell_text(r.cells[0], it, 9.5)
        cell_text(r.cells[1], "", 10)
        cell_text(r.cells[2], "", 9.5)
    style_table(t, [8.0, 2.4, 7.4])
    para(doc, "", 6, space_after=2)

check_table("説明パート（15分）", [
    "冒頭で「本日のゴール」を示した",
    "結論から話している（理由が先になっていない）",
    "数値・出典を資料のとおりに言えている",
    "いま資料のどこを話しているかが相手に分かる",
    "声の大きさ・スピード、顔を上げるタイミング",
    "時間内に収まっている",
])
check_table("実技パート（10分）", [
    "開始前に体調確認の声かけをした",
    "デモ → 参加者が実施 → 声かけ の流れになっている",
    "座ったままできる代替種目を提示した",
    "痛みが出た場合の中止指示を伝えた",
    "その講座のねらいから外れていない",
    "（オンラインの場合）カメラに全身が映る位置取り",
])
check_table("つなぎ", [
    "説明 → 実技",
    "実技 → 説明",
    "実技 → 習慣化シート",
])

# ---------------- シート4：ケース問答 ----------------
doc.add_page_break()
sheet_header(doc, "4", "ケース問答", "使用者：育成担当者　／　10分　※正解の暗記ではなく初動を確認する")

t = doc.add_table(rows=1, cols=3)
t.style = "Table Grid"
header_row(t, ["設問", "確認したい初動", "可／要補足"], [5.6, 9.4, 2.8], center_idx=(2,))
for q, a in [
    ("開始前に「腰が痛い」と申し出があった",
     "無理をしないよう伝え、座ったままの代替種目を案内。痛みが強ければ見学でも可と伝える。医療的な判断はしない"),
    ("開始5分前、オンラインに接続できない",
     "企業のご担当者に連絡し、同時に事務局へ電話。代替案（後日振替）は事務局が企業と調整する"),
    ("終了後、個別に治療の相談をされた",
     "共感の一言 → 医療者ではないと伝える → かかりつけ医・産業医への相談を勧める（固定フレーズ）"),
    ("企業から「後日パーソナルもお願いしたい」と言われた",
     "その場で受けない。営業・利益誘導は不可。事務局に持ち帰る"),
    ("実技で参加者がほとんど動いてくれない",
     "強制しない。座位でできる小さい動きに切り替え、人数ではなく一人ずつに声をかける"),
]:
    r = t.add_row()
    cell_text(r.cells[0], q, 9.5, bold=True)
    cell_text(r.cells[1], a, 9.5, color=SUB)
    cell_text(r.cells[2], "", 10)
style_table(t, [5.6, 9.4, 2.8])
para(doc, "", 8, space_after=2)
para(doc, "※ 5問すべてに共通する答えは「迷ったら事務局へ連絡する」。最後にこれを確認して終える。",
     9.5, color=SUB, space_after=0)

# ---------------- シート5：結果と補強プラン ----------------
doc.add_page_break()
sheet_header(doc, "5", "結果と補強プラン", "記入者：育成担当者　／　控えを本人にお渡しください")

hdr = doc.add_table(rows=2, cols=3)
hdr.style = "Table Grid"
header_row(hdr, ["受講者名", "チェック実施日", "担当講座"], [4.6, 4.6, 8.6])
for i, v in enumerate(["", "　　　年　　月　　日", ""]):
    cell_text(hdr.rows[1].cells[i], v, 10)
style_table(hdr, [4.6, 4.6, 8.6])
para(doc, "", 8, space_after=2)

para(doc, "結果", 12, bold=True, color=DEEP, space_after=4)
t = doc.add_table(rows=1, cols=3)
t.style = "Table Grid"
header_row(t, ["", "判定", "意味とその後の対応"], [1.6, 4.4, 11.8], center_idx=(0,))
for mark, name, desc in [
    ("□", "A. そのまま実施可", "過不足なし。単独で担当していただきます"),
    ("□", "B. 補足付きで実施可", "一部に補強が必要。補足資料・テキストの読み合わせを追加します"),
    ("□", "C. 同行サポート付き", "初回のみ同行・オンライン同席のうえ担当していただきます"),
]:
    r = t.add_row()
    cell_text(r.cells[0], mark, 12, align=WD_ALIGN_PARAGRAPH.CENTER)
    cell_text(r.cells[1], name, 10, bold=True, color=DEEP)
    cell_text(r.cells[2], desc, 9.5)
style_table(t, [1.6, 4.4, 11.8])
para(doc, "「担当していただかない」という結果はありません。", 9.5, bold=True, color=GREEN, space_after=8)

para(doc, "良かった点（2つ以上）", 11.5, bold=True, color=GREEN, space_after=3)
blank_box(doc, "", lines=2)
para(doc, "研修当日までに補強する点（3つまで）", 11.5, bold=True, color=ACCENT, space_after=3)
blank_box(doc, "", lines=3)
para(doc, "こちらで用意するもの（補足資料・読み合わせ・同行など）", 11.5, bold=True, color=GREEN, space_after=3)
blank_box(doc, "", lines=2)

para(doc, "", 6, space_after=2)
t = doc.add_table(rows=1, cols=3)
t.style = "Table Grid"
header_row(t, ["次のアクション", "期限", "担当"], [9.0, 4.4, 4.4])
for _ in range(3):
    r = t.add_row()
    for i in range(3):
        cell_text(r.cells[i], "", 10)
style_table(t, [9.0, 4.4, 4.4])
para(doc, "", 8, space_after=2)
para(doc, "ご署名　　　　　　　　　　　　　　　　　　　　　　　担当者確認　　　　　　　　　　　　　　　",
     10.5, space_after=0)

out = sys.argv[1] if len(sys.argv) > 1 else "docs/講師育成_60分チェック_実施セット.docx"
doc.save(out)
print("written:", out)
